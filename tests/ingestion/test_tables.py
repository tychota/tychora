"""Conservative table continuation decisions preserve original fragments."""

from tests.typed_json import json_array, json_object
from tychora.ingestion.tables import continuation_score


def test_missing_features_do_not_inflate_table_confidence():
    assert (
        continuation_score(
            headers=True,
            caption=False,
            heading=True,
            overlap=False,
            page_edges=False,
            marker=False,
        )
        == 0.5
    )


import pytest


@pytest.fixture
def fragments():
    from tychora.contracts.evidence import DocumentLocator, Geometry
    from tychora.contracts.tables import TableFragment

    def make(fragment_id, page, *, bbox):
        return TableFragment(
            fragment_id=fragment_id,
            revision_id="revision",
            extraction_id="extraction",
            unit_id=fragment_id + "-unit",
            native_ref="#/tables/" + fragment_id,
            page=page,
            column_count=2,
            grid=(("Sample", "Value"), ("A", "1")),
            header_rows=(0,),
            caption="Table 1: Samples",
            table_number="1",
            heading_path=("Observations",),
            heading_level=1,
            bbox=Geometry(
                domain="points", x0=bbox[0], y0=bbox[1], x1=bbox[2], y1=bbox[3]
            ),
            page_width=100.0,
            page_height=100.0,
            required_unit_ids=frozenset({fragment_id + "-unit"}),
            locators=(
                DocumentLocator(
                    artifact_ref="0" * 64,
                    native_ref="source.pdf",
                    char_start=0,
                    char_end=1,
                    physical_page=page,
                ),
            ),
        )

    return make("a", 1, bbox=(10, 80, 90, 99)), make("b", 2, bbox=(10, 1, 90, 20))


def test_high_confidence_continuation_is_derived_not_confirmed(fragments):
    from tychora.ingestion.tables import continuation

    a, b = fragments
    result = continuation(a, b)
    assert result.score == 0.9
    assert result.status == "derived"
    assert result.omit_repeated_header


@pytest.mark.parametrize(
    "change,reason",
    [
        ({"page": 3}, "nonadjacent_pages"),
        (
            {"column_count": 3, "grid": (("Sample", "Value", "Other"),)},
            "incompatible_columns",
        ),
        ({"table_number": "2"}, "different_table_number"),
        ({"intervening_heading_level": 1}, "intervening_heading"),
    ],
)
def test_hard_rejects_precede_confidence(fragments, change, reason):
    from tychora.ingestion.tables import continuation

    a, b = fragments
    result = continuation(a, b.model_copy(update=change))
    assert result.status == "separate" and result.score == 0
    assert reason in result.reasons


def test_same_headers_without_other_evidence_stay_separate(fragments):
    from tychora.ingestion.tables import continuation

    a, b = fragments
    a = a.model_copy(update={"caption": "", "table_number": None, "bbox": None})
    b = b.model_copy(update={"caption": "", "table_number": None, "bbox": None})
    result = continuation(a, b)
    assert result.score == 0.5 and result.status == "separate"


def test_missing_geometry_does_not_inflate_confidence(fragments):
    from tychora.ingestion.tables import continuation

    a, b = fragments
    result = continuation(a.model_copy(update={"bbox": None}), b)
    assert result.score == 0.7 and result.status == "proposed"


def test_reconstructed_headers_never_replace_original_fragments(fragments):
    from tychora.ingestion.tables import continuation, reconstruct_rows

    a, b = fragments
    rows = reconstruct_rows(a, b, continuation(a, b))
    assert [(row.fragment_id, row.row_index, row.cells) for row in rows] == [
        ("a", 0, ("Sample", "Value")),
        ("a", 1, ("A", "1")),
        ("b", 1, ("A", "1")),
    ]
    assert b.grid[0] == ("Sample", "Value") and b.header_rows == (0,)


@pytest.mark.parametrize(
    "profile,expected", [("documents", "derived"), ("verified", "confirmed_only")]
)
def test_resolved_continuation_policy_is_explicit(tmp_path, profile, expected):
    from tychora.configuration.resolve import resolve_config

    path = tmp_path / "workspace.toml"
    path.write_text(f'schema_version=1\n[profiles.test]\nbuiltin="{profile}@1"\n')
    resolved = resolve_config(path, "test", {})
    assert resolved.values.retrieval.table_continuations == expected
    assert (
        json_object(resolved.inspect(), "values", "retrieval")["table_continuations"]
        == expected
    )


def test_custom_continuation_policy_changes_query_not_extraction_identity(tmp_path):
    from tychora.configuration.resolve import resolve_config

    path = tmp_path / "workspace.toml"
    base = 'schema_version=1\n[profiles.test]\nbuiltin="documents@1"\n'
    path.write_text(base)
    derived = resolve_config(path, "test", {})
    path.write_text(base + '[configuration.retrieval]\ntable_continuations="none"\n')
    none = resolve_config(path, "test", {})
    assert none.values.retrieval.table_continuations == "none"
    assert derived.eligibility_fingerprint != none.eligibility_fingerprint
    assert derived.query_fingerprint != none.query_fingerprint
    assert derived.extraction_fingerprint == none.extraction_fingerprint


def test_continuation_assertion_has_complete_support_and_typed_jsonl(fragments):
    import json
    from hashlib import sha256

    from tests.helpers.corpus import artifact, registry
    from tychora.adapters.files.references import required_refs
    from tychora.contracts.evidence import EvidenceUnit
    from tychora.contracts.relationships import encode_assertions
    from tychora.ingestion.tables import (
        continuation,
        continuation_assertion,
        validate_assertion,
    )

    a, b = fragments
    units = {}
    for fragment in (a, b):
        text = "Sample\tValue\nA\t1"
        units[fragment.unit_id] = EvidenceUnit(
            unit_id=fragment.unit_id,
            extraction_id=fragment.extraction_id,
            parent_id=None,
            ordinal=len(units),
            unit_type="table",
            native_ref=fragment.native_ref,
            boundary_key=fragment.native_ref,
            text=text,
            text_sha=sha256(text.encode()).hexdigest(),
            locators=fragment.locators,
        )
    derivation = artifact(b'{"algorithm":"table-continuation-v1"}', "application/json")
    assertion = continuation_assertion(a, b, continuation(a, b), derivation)
    assert assertion is not None
    validate_assertion(
        assertion,
        units,
        registry(),
        proxy_ids=frozenset(),
        kind_id="manual@1",
        revision_id=a.revision_id,
    )
    assert assertion.required_support_ids == frozenset(units)
    assert assertion.method == "parser_inferred" and assertion.status == "recorded"
    assert assertion.confidence == 0.9
    records = [
        json.loads(line) for line in encode_assertions((assertion,)).splitlines()
    ]
    assert records[0]["type"] == "assertion" and records[0]["schema_version"] == "1"
    assert records[0]["assertion_id"] == assertion.assertion_id
    assert required_refs(records, capture_published=True) == frozenset(
        {derivation.sha256}
    )
    blocked = frozenset({a.unit_id})
    with pytest.raises(ValueError):
        validate_assertion(
            assertion,
            units,
            registry(),
            proxy_ids=blocked,
            kind_id="manual@1",
            revision_id=a.revision_id,
        )


@pytest.mark.parametrize(
    "policy,status,allowed",
    [
        ("derived", "recorded", True),
        ("derived", "proposed", False),
        ("derived", "rejected", False),
        ("confirmed_only", "recorded", False),
        ("confirmed_only", "confirmed", True),
        ("none", "confirmed", False),
    ],
)
def test_continuation_policy_only_narrows_eligible_relations(policy, status, allowed):
    from tychora.ingestion.tables import continuation_context_allowed

    def permitted(
        *,
        blocked: bool = False,
        eligible_original_ids: frozenset[str] = frozenset({"a", "b"}),
        permitted_predicates: frozenset[str] = frozenset({"continues_table@1"}),
    ) -> bool:
        return continuation_context_allowed(
            policy=policy,
            status=status,
            predicate="continues_table@1",
            permitted_predicates=permitted_predicates,
            permitted_statuses=frozenset({"recorded", "confirmed"}),
            required_support_ids=frozenset({"a", "b"}),
            eligible_original_ids=eligible_original_ids,
            blocked=blocked,
        )

    assert permitted() is allowed
    assert not permitted(blocked=True)
    assert not permitted(eligible_original_ids=frozenset({"a"}))
    assert not permitted(permitted_predicates=frozenset())


from tests.qualification import test_docling_records as qualification

document_record_case = qualification.document_record_case


def create_table_source(path, content):
    if path.suffix == ".pdf":
        from reportlab.lib import colors
        from reportlab.pdfgen import canvas
        from reportlab.platypus import Table, TableStyle

        pdf = canvas.Canvas(str(path), pagesize=(612, 792), invariant=1)
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(72, 710, content[0][1])
        pdf.setFont("Helvetica", 12)
        pdf.drawString(72, 656, "Requirement: Preserve the observed sample table.")
        for page, rows in enumerate(
            (
                [["Sample", "Value"], ["Alpha", "1"]],
                [["Sample", "Value"], ["Beta", "2"]],
            )
        ):
            if page:
                pdf.showPage()
            pdf.setFont("Helvetica", 12)
            pdf.drawString(
                72,
                173 if not page else 735,
                "Table 1. Samples" + (" continued" if page else ""),
            )
            table = Table(rows, colWidths=(190, 190), rowHeights=(32, 32))
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 1, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 12),
                    ]
                )
            )
            table.wrapOn(pdf, 380, 64)
            table.drawOn(pdf, 72, 100 if not page else 660)
        pdf.save()
    else:
        from docx import Document

        doc = Document()
        doc.add_heading(content[0][1], level=1)
        doc.add_paragraph("Requirement: Preserve the observed sample table.")
        for index, name in enumerate(("Alpha", "Beta")):
            if index:
                doc.add_page_break()
            doc.add_paragraph("Table 1. Samples", style="Caption")
            table = doc.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            for row, values in zip(
                table.rows, (("Sample", "Value"), (name, str(index + 1)))
            ):
                for cell, text in zip(row.cells, values):
                    cell.text = text
        doc.save(path)


@pytest.mark.qualification
@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_composed_tables_commit_once_with_exact_rows_and_captions(
    document_record_case, format
):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.composition import compose_document
    from tychora.ingestion.document_parts import validate_document_ownership

    source = document_record_case(
        "standard@1", format, source_builder=create_table_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    result = source.owner.run_blocking_sync(
        lambda: compose_document(
            source.parsed, source.registry, plan, source.owner.artifacts, source.ctx
        ),
        source.ctx,
    )
    bundle = result.bundle
    originals = tuple(
        unit
        for unit in bundle.units
        if unit.unit_id not in {p.proxy_unit_id for p in bundle.proxies}
    )
    validate_document_ownership(originals, source.parsed)
    assert len(result.fragments) == 2
    assert any(
        unit.unit_type == "table_row" and "Alpha" in unit.text for unit in originals
    )
    assert any(
        unit.unit_type == "table_row" and "Beta" in unit.text for unit in originals
    )
    assert any("Table 1. Samples" in unit.text for unit in originals)
    assert all(
        fragment.unit_id in fragment.required_unit_ids for fragment in result.fragments
    )
    assert all(len(fragment.required_unit_ids) > 1 for fragment in result.fragments)
    if format == "pdf":
        assert all(not fragment.caption for fragment in result.fragments)
        assert result.assertions and result.assertions[0].status in (
            "recorded",
            "proposed",
        )
    else:
        assert not result.assertions and all(
            fragment.page is None for fragment in result.fragments
        )
        assert not bundle.proxies and all(
            not fragment.caption for fragment in result.fragments
        )
    source.owner.run_blocking_sync(
        lambda: source.store.commit_extraction(bundle), source.ctx
    )
    source.owner.run_blocking_sync(
        lambda: source.store.commit_extraction(bundle), source.ctx
    )
    assert (
        source.owner.run_blocking_sync(
            lambda: source.owner.db.execute(
                "SELECT COUNT(*) FROM extractions"
            ).fetchone()[0],
            source.ctx,
        )
        == 1
    )
    assert (
        source.owner.run_blocking_sync(
            lambda: source.store.hydrate(tuple(u.unit_id for u in bundle.units)),
            source.ctx,
        )
        == bundle.units
    )


@pytest.mark.qualification
def test_composed_document_archive_replays_canonical_evidence(
    document_record_case, tmp_path, monkeypatch
):
    import hashlib
    import io
    import json
    import shutil
    import sqlite3
    from types import SimpleNamespace

    from tychora.adapters.docling.parser import DoclingParser
    from tychora.adapters.files.cas import LocalArtifacts
    from tychora.adapters.files.references import from_backup
    from tychora.application.backup import BackupManifest, BackupService
    from tychora.application.restore import RestoreService
    from tychora.configuration.registry import Registry
    from tychora.contracts.base import FrozenJsonMap
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.contracts.errors import TychoraError
    from tychora.contracts.identity import canonical_bytes
    from tychora.contracts.requests import RequestScope
    from tychora.ingestion.composition import compose_document
    from tychora.ports.execution import ExecutionContext

    scope = RequestScope(
        principal_id="owner",
        collection_ids=frozenset({"00000000-0000-4000-8000-000000000010"}),
        access_ceiling=3,
        capabilities=frozenset({"manage", "read", "write"}),
    )
    document = {
        "schema_name": "DoclingDocument",
        "version": "1.10.0",
        "name": "Synthetic",
        "body": {
            "self_ref": "#/body",
            "children": [
                {"$ref": "#/texts/0"},
                {"$ref": "#/texts/1"},
            ],
        },
        "furniture": {"self_ref": "#/furniture", "children": []},
        "texts": [
            {
                "self_ref": "#/texts/0",
                "parent": {"$ref": "#/body"},
                "children": [],
                "label": "section_header",
                "level": 1,
                "text": "STD-8 Sample",
                "orig": "STD-8 Sample",
                "prov": [],
            },
            {
                "self_ref": "#/texts/1",
                "parent": {"$ref": "#/body"},
                "children": [],
                "label": "text",
                "text": "Requirement: Keep the sample.",
                "orig": "Requirement: Keep the sample.",
                "prov": [],
            },
        ],
        "groups": [],
        "tables": [],
        "pictures": [],
        "pages": {},
    }

    def extract_synthetic(self, original, config, ctx):
        registry = Registry.from_config(config)
        kind = registry.kinds[self.registration.revision.kind_id]
        return self._retain(
            (
                (
                    SimpleNamespace(name="document.json"),
                    io.BytesIO(canonical_bytes(document)),
                ),
            ),
            original,
            config,
            registry,
            kind,
            FrozenJsonMap({}),
        )

    monkeypatch.setattr(DoclingParser, "extract", extract_synthetic)
    monkeypatch.setenv("TYCHORA_TEST_DOCLING_MODELS", str(tmp_path))
    source = document_record_case("standard@1", "docx", commit=False, scope=scope)
    result = source.owner.run_blocking_sync(
        lambda: compose_document(
            source.parsed,
            source.registry,
            DocumentCompositionPlan(
                base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
            ),
            source.owner.artifacts,
            source.ctx,
        ),
        source.ctx,
    )
    source.owner.run_blocking_sync(
        lambda: source.store.commit_extraction(result.bundle), source.ctx
    )
    archive = tmp_path / "archive"
    BackupService(source.owner).export(archive, scope, ctx=ExecutionContext())
    RestoreService().import_backup(
        archive, tmp_path / "restored", scope, ExecutionContext()
    )

    stripped = tmp_path / "stripped-mapping"
    shutil.copytree(archive, stripped)
    stripped_artifacts = LocalArtifacts(stripped / "artifacts")
    stripped_db = sqlite3.connect(stripped / "corpus.sqlite3")
    try:
        composed_sha = stripped_db.execute(
            "SELECT document_artifact_sha FROM extractions"
        ).fetchone()[0]
        composed = json.loads((stripped / "artifacts" / composed_sha).read_bytes())
        raw_link = composed["required_manifests"][0]
        raw_sha = raw_link["artifact"]["sha256"]
        raw = json.loads((stripped / "artifacts" / raw_sha).read_bytes())
        assert raw["data"].pop("docling_json_ref")
        assert raw["data"]["registration"]["format"] == "docx"
        raw["data"]["registration"]["format"] = "text"
        stripped_db.execute(
            "UPDATE locators SET locator_type='text' WHERE locator_type='document'"
        )

        def replace_artifact(old_sha, value):
            media = stripped_db.execute(
                "SELECT media_type FROM artifacts WHERE artifact_sha=?", (old_sha,)
            ).fetchone()[0]
            ref = stripped_artifacts.put_stream(
                io.BytesIO(canonical_bytes(value)), media
            )
            stripped_db.execute(
                "INSERT INTO artifacts (artifact_sha,size_bytes,media_type,storage_key,created_at) "
                "SELECT ?,?,?,?,created_at FROM artifacts WHERE artifact_sha=?",
                (ref.sha256, ref.size_bytes, ref.media_type, ref.storage_key, old_sha),
            )
            return ref

        new_raw = replace_artifact(raw_sha, raw)
        raw_link["artifact"] = new_raw.model_dump(mode="json")
        new_composed = replace_artifact(composed_sha, composed)
        stripped_db.execute(
            "UPDATE extractions SET document_artifact_sha=?", (new_composed.sha256,)
        )
        stripped_db.commit()
        inventory = from_backup(stripped_db, artifacts=stripped_artifacts)
    finally:
        stripped_db.close()
    old_manifest = json.loads((stripped / "backup-manifest.json").read_bytes())
    new_manifest = BackupManifest(
        workspace_id=inventory.workspace_id,
        artifacts=inventory.artifacts,
        payload_states=inventory.payload_states,
        policy_epoch=inventory.policy_epoch,
        lexical_statistics_epoch=inventory.lexical_statistics_epoch,
        active_snapshot_id=inventory.active_snapshot_id,
        snapshot_ids=inventory.snapshot_ids,
        tombstones=inventory.tombstones,
        config_shas=inventory.config_shas,
        resolved_config_sha256=old_manifest["data"]["resolved_config_sha256"],
    ).as_header()
    (stripped / "backup-manifest.json").write_bytes(canonical_bytes(new_manifest))
    stripped_receipt_path = stripped / "backup-receipt.json"
    stripped_receipt = json.loads(stripped_receipt_path.read_bytes())
    stripped_receipt["sqlite_sha256"] = hashlib.sha256(
        (stripped / "corpus.sqlite3").read_bytes()
    ).hexdigest()
    stripped_receipt["manifest_sha256"] = hashlib.sha256(
        (stripped / "backup-manifest.json").read_bytes()
    ).hexdigest()
    stripped_receipt_path.write_bytes(canonical_bytes(stripped_receipt))
    with pytest.raises(TychoraError, match="Retained document source mapping differs"):
        RestoreService().import_backup(
            stripped, tmp_path / "stripped-restored", scope, ExecutionContext()
        )
    assert not (tmp_path / "stripped-restored").exists()

    missing_marker = tmp_path / "missing-marker"
    shutil.copytree(archive, missing_marker)
    marker_db = sqlite3.connect(missing_marker / "corpus.sqlite3")
    try:
        marker_db.execute(
            "UPDATE evidence_units SET attributes_json="
            "json_remove(attributes_json,'$.document_representation') "
            "WHERE text_origin='source'"
        )
        marker_db.commit()
    finally:
        marker_db.close()
    marker_receipt_path = missing_marker / "backup-receipt.json"
    marker_receipt = json.loads(marker_receipt_path.read_bytes())
    marker_receipt["sqlite_sha256"] = hashlib.sha256(
        (missing_marker / "corpus.sqlite3").read_bytes()
    ).hexdigest()
    marker_receipt_path.write_bytes(canonical_bytes(marker_receipt))
    with pytest.raises(TychoraError, match="Retained document source mapping differs"):
        RestoreService().import_backup(
            missing_marker, tmp_path / "marker-restored", scope, ExecutionContext()
        )
    assert not (tmp_path / "marker-restored").exists()

    db = sqlite3.connect(archive / "corpus.sqlite3")
    try:
        db.execute(
            "UPDATE evidence_units SET text=? WHERE unit_id=?",
            ("altered canonical evidence", result.bundle.units[0].unit_id),
        )
        db.commit()
    finally:
        db.close()
    receipt_path = archive / "backup-receipt.json"
    receipt = json.loads(receipt_path.read_bytes())
    receipt["sqlite_sha256"] = hashlib.sha256(
        (archive / "corpus.sqlite3").read_bytes()
    ).hexdigest()
    receipt_path.write_bytes(canonical_bytes(receipt))
    with pytest.raises(TychoraError, match="Retained document source mapping differs"):
        RestoreService().import_backup(
            archive, tmp_path / "altered", scope, ExecutionContext()
        )


def composed_inputs(source):
    def read():
        values = {}
        for asset in source.parsed.assets:
            with source.owner.artifacts.open_verified(asset.artifact_ref) as stream:
                values[asset.artifact_ref.sha256] = stream.read()
        return values

    return source.owner.run_blocking_sync(read, source.ctx)


@pytest.mark.qualification
def test_composed_replay_requires_retained_derivations(document_record_case):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.composition import (
        build_composed_document,
        replay_composed_document,
    )

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_table_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    inputs = composed_inputs(source)
    result = build_composed_document(source.parsed, source.registry, plan, inputs)
    retained = {**inputs, **result.payloads}
    assert (
        replay_composed_document(
            result.payloads[result.manifest_ref.sha256],
            result.manifest_ref,
            source.parsed,
            source.registry,
            retained,
        ).bundle
        == result.bundle
    )
    derivation = next(
        sha for sha in result.payloads if sha != result.manifest_ref.sha256
    )
    retained.pop(derivation)
    with pytest.raises(ValueError, match="retained"):
        replay_composed_document(
            result.payloads[result.manifest_ref.sha256],
            result.manifest_ref,
            source.parsed,
            source.registry,
            retained,
        )


@pytest.mark.qualification
def test_composition_plan_and_manifest_forgery_fail_before_commit(document_record_case):
    import json

    from tests.helpers.corpus import artifact
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.contracts.identity import canonical_bytes
    from tychora.ingestion.composition import (
        build_composed_document,
        replay_composed_document,
    )

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_table_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    inputs = composed_inputs(source)
    result = build_composed_document(source.parsed, source.registry, plan, inputs)
    with pytest.raises(ValueError, match="base normalizer"):
        build_composed_document(
            source.parsed,
            source.registry,
            plan.model_copy(update={"base_normalizer_fingerprint": "0" * 64}),
            inputs,
        )
    root = json.loads(result.payloads[result.manifest_ref.sha256])
    root["data"]["plan"]["caption_copies"] = False
    forged = canonical_bytes(root)
    with pytest.raises(ValueError, match="replay inputs"):
        replay_composed_document(
            forged,
            artifact(forged, "application/json"),
            source.parsed,
            source.registry,
            {**inputs, **result.payloads},
        )


@pytest.mark.qualification
def test_changed_plan_commits_distinct_extraction_same_revision(document_record_case):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.composition import compose_document

    source = document_record_case(
        "standard@1", "docx", source_builder=create_table_source, commit=False
    )
    plans = [
        DocumentCompositionPlan(
            base_normalizer_fingerprint=source.parsed.normalizer_fingerprint,
            caption_copies=copies,
        )
        for copies in (True, False)
    ]
    products = [
        source.owner.run_blocking_sync(
            lambda plan=plan: compose_document(
                source.parsed, source.registry, plan, source.owner.artifacts, source.ctx
            ),
            source.ctx,
        )
        for plan in plans
    ]
    assert products[0].bundle.revision == products[1].bundle.revision
    assert (
        products[0].bundle.extraction.extraction_id
        != products[1].bundle.extraction.extraction_id
    )
    for product in products:
        source.owner.run_blocking_sync(
            lambda product=product: source.store.commit_extraction(product.bundle),
            source.ctx,
        )
    assert (
        source.owner.run_blocking_sync(
            lambda: source.owner.db.execute(
                "SELECT COUNT(*) FROM extractions"
            ).fetchone()[0],
            source.ctx,
        )
        == 2
    )


def create_rich_source(path, content):
    import shutil
    from pathlib import Path

    shutil.copyfile(Path(__file__).parents[1] / "fixtures/documents/rich.pdf", path)


@pytest.mark.qualification
def test_real_caption_copy_commit_and_retention_replay(document_record_case):
    import json

    from tychora.adapters.files.references import required_closure
    from tychora.contracts.artifacts import ManifestRef
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.composition import compose_document, replay_composed_document
    from tychora.ingestion.documents import replay_document

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_rich_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    result = source.owner.run_blocking_sync(
        lambda: compose_document(
            source.parsed, source.registry, plan, source.owner.artifacts, source.ctx
        ),
        source.ctx,
    )
    assert result.bundle.proxies
    caption_copy = next(
        proxy for proxy in result.bundle.proxies if proxy.origin_type == "caption_copy"
    )
    units = {unit.unit_id: unit for unit in result.bundle.units}
    assert not units[caption_copy.proxy_unit_id].citable
    assert any(
        units[uid].citable and "Reference vessel" in units[uid].text
        for uid in caption_copy.required_unit_ids
    )
    source.owner.run_blocking_sync(
        lambda: source.store.commit_extraction(result.bundle), source.ctx
    )

    def retained_replay():
        def payload(ref):
            assert ref.sha256 != source.parsed.original_ref.sha256
            with source.owner.artifacts.open_verified(ref) as stream:
                return stream.read()

        def load(link):
            value = payload(link.artifact)
            return (
                [json.loads(line) for line in value.splitlines()]
                if link.type == "assertion"
                else json.loads(value)
            )

        closure = required_closure(
            (
                ManifestRef(
                    artifact=result.manifest_ref, type="extraction", schema_version="1"
                ),
            ),
            load=load,
        )
        assert source.parsed.original_ref.sha256 not in closure
        raw = replay_document(
            payload(source.parsed.raw_document_ref),
            payload(source.parsed.docling_json_ref),
            source.parsed.raw_document_ref,
        )
        refs = {ref.sha256: ref for ref in result.bundle.artifacts}
        inputs = {sha: payload(refs[sha]) for sha in closure if sha in refs}
        replay = replay_composed_document(
            payload(result.manifest_ref),
            result.manifest_ref,
            raw,
            source.registry,
            inputs,
        )
        assert replay.bundle == result.bundle

    source.owner.run_blocking_sync(retained_replay, source.ctx)


@pytest.mark.qualification
def test_composition_metadata_observations_do_not_change_semantic_identity(
    document_record_case,
):
    from tests.ingestion.test_document_records import rebind
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.document_identity import DocumentNormalizationContext

    source = document_record_case(
        "standard@1", "docx", source_builder=create_table_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    data = source.parsed.model_dump(mode="json", exclude={"raw_document_ref"})
    data["registration"]["revision"]["captured_at"] = "2026-09-12T00:00:00+00:00"
    data["registration"]["capture"]["captured_at"] = "2026-09-12T00:00:00Z"
    data["registration"]["source"]["created_at"] = "2026-09-12T00:00:00+00:00"
    later = rebind(data)
    before = DocumentNormalizationContext.bind(source.parsed, source.registry, plan)
    after = DocumentNormalizationContext.bind(later, source.registry, plan)
    assert before.raw_manifest_sha != after.raw_manifest_sha
    assert before.normalizer_fingerprint == after.normalizer_fingerprint
    with pytest.raises(ValueError, match="context"):
        before.verify(later, source.registry)


@pytest.mark.qualification
def test_supplied_proxy_plan_binds_text_derivation_and_exact_original_selectors(
    document_record_case,
):
    from hashlib import sha256

    from tests.helpers.corpus import artifact
    from tychora.contracts.composition import (
        DocumentCompositionPlan,
        PlannedProxy,
        SourceSelector,
    )
    from tychora.ingestion.composition import build_composed_document

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_rich_source, commit=False
    )
    asset = next(
        asset for asset in source.parsed.assets if asset.asset_type == "picture"
    )
    text = "A blue marker appears beside a vessel."
    selector = SourceSelector(native_ref=asset.native_ref)
    receipt = artifact(b'{"method":"fixture visual description"}', "application/json")
    proxy = PlannedProxy(
        asset_native_ref=asset.native_ref,
        kind="description",
        text=text,
        text_sha=sha256(text.encode()).hexdigest(),
        dependencies=(selector,),
        derivation=receipt,
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint,
        proxies=(proxy,),
    )
    inputs = {
        **composed_inputs(source),
        receipt.sha256: b'{"method":"fixture visual description"}',
    }
    result = build_composed_document(source.parsed, source.registry, plan, inputs)
    generated = next(unit for unit in result.bundle.units if unit.text == text)
    assert not generated.citable and generated.review_state == "needs_review"
    assert generated.locators == (asset.locator,)
    with pytest.raises(ValueError, match="payload"):
        build_composed_document(
            source.parsed, source.registry, plan, {**inputs, receipt.sha256: b"forged"}
        )
    wrong = proxy.model_copy(
        update={"dependencies": (SourceSelector(native_ref="#/pictures/missing"),)}
    )
    with pytest.raises(ValueError, match="SDK item"):
        build_composed_document(
            source.parsed,
            source.registry,
            plan.model_copy(update={"proxies": (wrong,)}),
            inputs,
        )
    receipt2 = artifact(b"different recorded derivation", "text/plain")
    changed = plan.model_copy(
        update={"proxies": (proxy.model_copy(update={"derivation": receipt2}),)}
    )
    other = build_composed_document(
        source.parsed,
        source.registry,
        changed,
        {**inputs, receipt2.sha256: b"different recorded derivation"},
    )
    assert (
        result.bundle.extraction.extraction_id != other.bundle.extraction.extraction_id
    )


def test_fragment_support_serialization_has_stable_order(fragments):
    a, _ = fragments
    support = frozenset({a.unit_id, *(f"owner-{i}" for i in range(20))})
    assert a.model_copy(update={"required_unit_ids": support}).model_dump(mode="json")[
        "required_unit_ids"
    ] == sorted(support)


@pytest.mark.qualification
def test_assertion_cannot_drop_overlapping_original_owners(document_record_case):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.contracts.relationships import (
        AssertionObservation,
        RelationshipAssertion,
    )
    from tychora.ingestion.composition import build_composed_document
    from tychora.ingestion.tables import validate_assertion

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_table_source, commit=False
    )
    result = build_composed_document(
        source.parsed,
        source.registry,
        DocumentCompositionPlan(
            base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
        ),
        composed_inputs(source),
    )
    assertion = result.assertions[0]
    data = assertion.model_dump(exclude={"assertion_id"})
    data["required_support_ids"] = frozenset(
        {assertion.subject.unit_id, assertion.object.unit_id, assertion.origin_unit_id}
    )
    forged = RelationshipAssertion.from_observation(
        AssertionObservation.model_validate(data)
    )
    assert forged.required_support_ids < assertion.required_support_ids
    with pytest.raises(ValueError, match="overlapping"):
        validate_assertion(
            forged,
            {unit.unit_id: unit for unit in result.bundle.units},
            source.registry,
            proxy_ids=frozenset(),
            kind_id=result.bundle.revision.kind_id,
            revision_id=result.bundle.revision.revision_id,
        )


@pytest.mark.qualification
def test_composed_manifest_capacity_exact_boundary_and_work_limit(
    document_record_case, monkeypatch
):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.contracts.errors import TychoraError
    from tychora.ingestion import composition, document_parts

    source = document_record_case(
        "standard@1", "docx", source_builder=create_table_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    inputs = composed_inputs(source)
    result = composition.build_composed_document(
        source.parsed, source.registry, plan, inputs
    )
    with monkeypatch.context() as patch:
        patch.setattr(
            composition, "MAX_COMPOSED_MANIFEST_BYTES", result.manifest_ref.size_bytes
        )
        assert (
            composition.build_composed_document(
                source.parsed, source.registry, plan, inputs
            ).bundle
            == result.bundle
        )
        patch.setattr(
            composition,
            "MAX_COMPOSED_MANIFEST_BYTES",
            result.manifest_ref.size_bytes - 1,
        )
        with pytest.raises(ValueError, match="manifest exceeds capacity"):
            composition.build_composed_document(
                source.parsed, source.registry, plan, inputs
            )
    monkeypatch.setattr(document_parts, "MAX_SEMANTIC_WORK", 1)
    with pytest.raises(TychoraError, match="SEMANTIC_EXPANSION_LIMIT"):
        composition.build_composed_document(
            source.parsed, source.registry, plan, inputs
        )


@pytest.mark.qualification
@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_grouped_provision_table_owners_remain_complete(document_record_case, format):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.composition import build_composed_document
    from tychora.ingestion.document_parts import validate_document_ownership

    source = document_record_case(
        "law@1", format, source_builder=create_table_source, commit=False
    )
    result = build_composed_document(
        source.parsed,
        source.registry,
        DocumentCompositionPlan(
            base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
        ),
        composed_inputs(source),
    )
    validate_document_ownership(result.bundle.units, source.parsed)
    table = result.fragments[0]
    owners = {unit.unit_id: unit for unit in result.bundle.units}
    assert any(
        owners[uid].native_ref.startswith("docling-record:numbered_provisions@1:")
        for uid in table.required_unit_ids
    )
    assert table.source_cells


@pytest.mark.parametrize(
    "forgery", ["predicate", "span", "cross_extraction", "endpoint_type"]
)
def test_assertion_rejects_invalid_registered_grounding(fragments, forgery):
    from hashlib import sha256

    from tests.helpers.corpus import artifact, registry
    from tychora.contracts.evidence import EvidenceUnit
    from tychora.contracts.relationships import (
        AssertionObservation,
        RelationshipAssertion,
        SupportSpan,
    )
    from tychora.ingestion.tables import (
        continuation,
        continuation_assertion,
        validate_assertion,
    )

    a, b = fragments
    units = {
        fragment.unit_id: EvidenceUnit(
            unit_id=fragment.unit_id,
            extraction_id=fragment.extraction_id,
            parent_id=None,
            ordinal=i,
            unit_type="table",
            native_ref=fragment.native_ref,
            boundary_key=fragment.native_ref,
            text="Sample",
            text_sha=sha256(b"Sample").hexdigest(),
            locators=fragment.locators,
        )
        for i, fragment in enumerate(fragments)
    }
    assertion = continuation_assertion(
        a, b, continuation(a, b), artifact(b"receipt", "text/plain")
    )
    assert assertion is not None
    data = assertion.model_dump(exclude={"assertion_id"})
    if forgery == "predicate":
        data["predicate"] = "unknown_relation@1"
    elif forgery == "span":
        data["support_spans"] = (
            SupportSpan(unit_id=a.unit_id, char_start=0, char_end=7),
        )
    elif forgery == "cross_extraction":
        units[b.unit_id] = units[b.unit_id].model_copy(
            update={"extraction_id": "other-extraction"}
        )
    else:
        units[b.unit_id] = units[b.unit_id].model_copy(
            update={"unit_type": "paragraph"}
        )
    forged = RelationshipAssertion.from_observation(
        AssertionObservation.model_validate(data)
    )
    with pytest.raises(ValueError):
        validate_assertion(
            forged,
            units,
            registry(),
            proxy_ids=frozenset(),
            kind_id="manual@1",
            revision_id=a.revision_id,
        )


def test_assertion_observations_remain_distinct_and_duplicates_fail(fragments):
    from tests.helpers.corpus import artifact
    from tychora.contracts.relationships import (
        AssertionObservation,
        RelationshipAssertion,
        encode_assertions,
    )
    from tychora.ingestion.tables import continuation, continuation_assertion

    a, b = fragments
    first = continuation_assertion(
        a, b, continuation(a, b), artifact(b"receipt", "text/plain")
    )
    assert first is not None
    data = first.model_dump(exclude={"assertion_id"})
    data["observation_key"] = "independent-observation"
    second = RelationshipAssertion.from_observation(
        AssertionObservation.model_validate(data)
    )
    assert first.assertion_id != second.assertion_id
    assert len(encode_assertions((first, second)).splitlines()) == 2
    with pytest.raises(ValueError, match="Duplicate"):
        encode_assertions((first, first))


@pytest.mark.qualification
def test_invalid_composition_plan_fails_before_artifact_reads(document_record_case):
    from contextlib import AbstractContextManager
    from typing import BinaryIO

    from tychora.contracts.artifacts import ArtifactRef, ArtifactStat
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion.composition import compose_document

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_rich_source, commit=False
    )

    class ForbiddenArtifacts:
        def open_verified(self, ref: ArtifactRef) -> AbstractContextManager[BinaryIO]:
            pytest.fail("Invalid composition opened retained payloads before binding")

        def put_stream(self, stream: BinaryIO, media_type: str) -> ArtifactRef:
            pytest.fail("Invalid composition wrote an artifact before binding")

        def stat(self, ref: ArtifactRef) -> ArtifactStat:
            pytest.fail("Invalid composition inspected an artifact before binding")

        def pin_payloads(self, shas: frozenset[str]) -> AbstractContextManager[None]:
            pytest.fail("Invalid composition pinned payloads before binding")

        def payload_guard(self, sha: str) -> AbstractContextManager[None]:
            pytest.fail("Invalid composition guarded an artifact before binding")

        def payload_is_pinned(self, sha: str) -> bool:
            pytest.fail("Invalid composition queried pin state before binding")

    with pytest.raises(ValueError, match="base normalizer"):
        compose_document(
            source.parsed,
            source.registry,
            DocumentCompositionPlan(base_normalizer_fingerprint="0" * 64),
            ForbiddenArtifacts(),
            source.ctx,
        )


@pytest.mark.parametrize(
    "root",
    [
        [],
        {"type": "extraction", "schema_version": "1", "data": []},
        {
            "type": "extraction",
            "schema_version": "1",
            "data": {"representation": "document-composition@1"},
        },
    ],
)
def test_composed_replay_rejects_malformed_envelopes(root):
    from tests.helpers.corpus import artifact
    from tychora.contracts.identity import canonical_bytes
    from tychora.ingestion.composition import replay_composed_document

    manifest = canonical_bytes(root)
    with pytest.raises(ValueError):
        # Malformed bytes must be rejected before the deliberately absent source context.
        replay_composed_document(
            manifest,
            artifact(manifest, "application/json"),
            None,  # ty: ignore[invalid-argument-type]
            None,  # ty: ignore[invalid-argument-type]
            {},
        )


@pytest.mark.qualification
def test_source_view_expansion_accounts_existing_original_parts(
    document_record_case, monkeypatch
):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.contracts.errors import TychoraError
    from tychora.ingestion import document_parts
    from tychora.ingestion.document_identity import DocumentNormalizationContext
    from tychora.ingestion.document_records import normalize_document_records
    from tychora.ingestion.document_views import source_views

    source = document_record_case("standard@1", "docx", commit=False)
    context = DocumentNormalizationContext.bind(
        source.parsed,
        source.registry,
        DocumentCompositionPlan(
            base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
        ),
    )
    base = normalize_document_records(
        source.parsed,
        source.parsed.registration.revision.revision_id,
        source.registry,
        context=context,
    )
    required = sum(
        1
        + len(json_array(unit.attributes["source_parts"]))
        + len(json_array(unit.attributes["structural_separators"]))
        for unit in base.units
    )
    monkeypatch.setattr(document_parts, "MAX_SEMANTIC_WORK", required - 1)
    with pytest.raises(TychoraError, match="SEMANTIC_EXPANSION_LIMIT"):
        source_views(source.parsed, context, base)


@pytest.mark.qualification
def test_proxy_expansion_stops_before_constructing_oversized_whole_result(
    document_record_case, monkeypatch
):
    from hashlib import sha256

    from tests.helpers.corpus import artifact
    from tychora.contracts.composition import (
        DocumentCompositionPlan,
        PlannedProxy,
        SourceSelector,
    )
    from tychora.contracts.errors import TychoraError
    from tychora.contracts.identity import canonical_bytes
    from tychora.ingestion import document_parts
    from tychora.ingestion.assets import ProxyFactory
    from tychora.ingestion.composition import build_composed_document

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_rich_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    inputs = composed_inputs(source)
    baseline = build_composed_document(source.parsed, source.registry, plan, inputs)
    ceiling = 3 * len(canonical_bytes(baseline.bundle.model_dump(mode="json")))
    asset = next(
        asset for asset in source.parsed.assets if asset.asset_type == "picture"
    )
    receipt = artifact(b"bounded derived input", "text/plain")
    planned = []
    for i in range(8):
        text = str(i) + "x" * (ceiling // 2)
        planned.append(
            PlannedProxy(
                asset_native_ref=asset.native_ref,
                kind="description",
                text=text,
                text_sha=sha256(text.encode()).hexdigest(),
                dependencies=(SourceSelector(native_ref=asset.native_ref),),
                derivation=receipt,
            )
        )
    calls = 0
    original = ProxyFactory.make_proxy

    def counted(self, asset_id, kind, text, dependencies, derivation):
        nonlocal calls
        if kind == "description":
            calls += 1
        return original(self, asset_id, kind, text, dependencies, derivation)

    monkeypatch.setattr(ProxyFactory, "make_proxy", counted)
    monkeypatch.setattr(document_parts, "MAX_SEMANTIC_BYTES", ceiling)
    with pytest.raises(TychoraError, match="SEMANTIC_EXPANSION_LIMIT"):
        build_composed_document(
            source.parsed,
            source.registry,
            plan.model_copy(update={"proxies": tuple(planned)}),
            {**inputs, receipt.sha256: b"bounded derived input"},
        )
    assert 0 < calls < len(planned)


@pytest.mark.qualification
def test_build_at_exact_input_ceiling_replays_retained_closure(
    document_record_case, monkeypatch
):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion import composition

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_table_source, commit=False
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
    )
    inputs = composed_inputs(source)
    ceiling = sum(len(payload) for payload in inputs.values())
    monkeypatch.setattr(composition, "MAX_COMPOSITION_INPUT_BYTES", ceiling)
    result = composition.build_composed_document(
        source.parsed, source.registry, plan, inputs
    )
    retained = {**inputs, **result.payloads}

    def raw_payload():
        with source.owner.artifacts.open_verified(
            source.parsed.docling_json_ref
        ) as stream:
            return stream.read()

    retained[source.parsed.docling_json_ref.sha256] = source.owner.run_blocking_sync(
        raw_payload, source.ctx
    )
    assert (
        composition.replay_composed_document(
            result.payloads[result.manifest_ref.sha256],
            result.manifest_ref,
            source.parsed,
            source.registry,
            retained,
        ).bundle
        == result.bundle
    )
    assert (
        source.owner.run_blocking_sync(
            lambda: composition.compose_document(
                source.parsed, source.registry, plan, source.owner.artifacts, source.ctx
            ),
            source.ctx,
        ).bundle
        == result.bundle
    )
    monkeypatch.setattr(composition, "MAX_COMPOSITION_INPUT_BYTES", ceiling - 1)
    with pytest.raises(ValueError, match="inputs exceed"):
        composition.build_composed_document(
            source.parsed, source.registry, plan, retained
        )
    with pytest.raises(ValueError, match="inputs exceed"):
        composition.replay_composed_document(
            result.payloads[result.manifest_ref.sha256],
            result.manifest_ref,
            source.parsed,
            source.registry,
            retained,
        )


@pytest.mark.qualification
def test_replay_has_its_own_exact_required_payload_bound(
    document_record_case, monkeypatch
):
    from tychora.contracts.composition import DocumentCompositionPlan
    from tychora.ingestion import composition

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_table_source, commit=False
    )
    result = composition.build_composed_document(
        source.parsed,
        source.registry,
        DocumentCompositionPlan(
            base_normalizer_fingerprint=source.parsed.normalizer_fingerprint
        ),
        composed_inputs(source),
    )
    retained = {**composed_inputs(source), **result.payloads}
    required_size = sum(len(payload) for payload in retained.values())
    monkeypatch.setattr(composition, "MAX_COMPOSITION_REPLAY_BYTES", required_size)
    assert (
        composition.replay_composed_document(
            result.payloads[result.manifest_ref.sha256],
            result.manifest_ref,
            source.parsed,
            source.registry,
            retained,
        ).bundle
        == result.bundle
    )
    monkeypatch.setattr(composition, "MAX_COMPOSITION_REPLAY_BYTES", required_size - 1)
    with pytest.raises(ValueError, match="replay.*capacity"):
        composition.replay_composed_document(
            result.payloads[result.manifest_ref.sha256],
            result.manifest_ref,
            source.parsed,
            source.registry,
            retained,
        )


@pytest.mark.qualification
def test_required_build_inputs_reject_conflicting_duplicate_refs(document_record_case):
    from hashlib import sha256

    from tychora.contracts.composition import (
        DocumentCompositionPlan,
        PlannedProxy,
        SourceSelector,
    )
    from tychora.ingestion.composition import build_composed_document

    source = document_record_case(
        "standard@1", "pdf", source_builder=create_table_source, commit=False
    )
    ref = source.parsed.assets[0].artifact_ref
    proxy = PlannedProxy(
        asset_native_ref=source.parsed.assets[0].native_ref,
        kind="description",
        text="Observed",
        text_sha=sha256(b"Observed").hexdigest(),
        dependencies=(SourceSelector(native_ref=source.parsed.units[0].native_ref),),
        derivation=ref.model_copy(update={"size_bytes": ref.size_bytes + 1}),
    )
    plan = DocumentCompositionPlan(
        base_normalizer_fingerprint=source.parsed.normalizer_fingerprint,
        proxies=(proxy,),
    )
    with pytest.raises(ValueError, match="Conflicting"):
        build_composed_document(
            source.parsed, source.registry, plan, composed_inputs(source)
        )
