"""Self-authored source-native semantic documents (no parser output fixtures)."""

from pathlib import Path


def record_content(kind: str) -> list[tuple[int, str]]:
    if kind == "standard@1":
        return [
            (1, "STD-01 - Sample trace"),
            (0, "Requirement: Record the sample before changing the instrument."),
            (0, "Keep the original label."),
            (0, "Expected evidence: The sample identifier is quotable."),
            (0, "Control: Inspect the source."),
            (0, "Severity: warning"),
            (0, "Foundation: REF-MISSING"),
        ]
    if kind == "template@1":
        return [
            (1, "Summary"),
            (0, "Instruction: State the observed result."),
            (0, "Keep the uncertainty explicit."),
            (1, "Limits"),
            (0, "Instruction: State the supported range."),
        ]
    if kind.startswith("feedback"):
        return [
            (1, "Observation"),
            (0, "Comment: The baseline is missing."),
            (0, "Target: paragraph-3"),
            *(
                [
                    (0, "Author: Source reviewer"),
                    (0, "Decision: Request source comparison"),
                    (0, "Observation: Source supplied observation"),
                    (0, "Status: proposed"),
                ]
                if kind == "feedback-explicit@1"
                else []
            ),
        ]
    if kind == "law@1":
        return [
            (1, "Article 4"),
            (0, "1. Record the sample."),
            (0, "2. Preserve the label."),
            (2, "4.1 Observations"),
            (0, "1. Retain the observation."),
            (1, "Article 5"),
            (0, "1. Inspect the source."),
        ]
    raise ValueError(kind)


def create_record_document(
    path: Path, content: list[tuple[int, str]], *, page_break_at: int | None = None
) -> None:
    if path.suffix == ".docx":
        from docx import Document

        doc = Document()
        for index, (level, text) in enumerate(content):
            if index == page_break_at:
                doc.add_page_break()
            if level:
                doc.add_heading(text, level=level)
            else:
                doc.add_paragraph(text)
        doc.save(str(path))
    elif path.suffix == ".pdf":
        from xml.sax.saxutils import escape

        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        flow = []
        for index, (level, text) in enumerate(content):
            if index == page_break_at:
                flow.append(PageBreak())
            flow.append(
                Paragraph(
                    escape(text), styles[f"Heading{level}" if level else "BodyText"]
                )
            )
            flow.append(Spacer(1, 48))
        SimpleDocTemplate(str(path), pagesize=(612, 792)).build(flow)
    else:
        path.write_text(
            "\n\n".join(
                "#" * level + " " + text if level else text for level, text in content
            )
            + "\n"
        )
