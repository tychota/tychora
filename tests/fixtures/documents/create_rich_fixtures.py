"""Generate original MIT-licensed synthetic fixtures; development use only.

Run with the locked test-fixtures dependency group. Coordinates in the manifest
are manually reviewed source drawing regions in PDF points, top-left origin.

The PDF embeds subsets of Bitstream Vera Sans (shipped with reportlab, under the
Bitstream Vera font licence, which permits embedding). With the standard
Helvetica left unembedded, Docling's page renderer substituted whatever
sans-serif the host had: Liberation Sans on the Ubuntu runners produced a page
on which the layout model read the figure caption as body text, so no caption
was linked to the figure. Embedded glyphs render the same on every host.
"""

from pathlib import Path

PASSAGES = (
    "Field notebook",
    "Record the sample before changing the instrument.",
    "Keep the blue marker beside the reference vessel.",
    "Repeat the observation after the indicator settles.",
    "Figure 1. Reference vessel and blue marker.",
)


def build(destination: Path) -> None:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image, ImageDraw
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("Vera", "Vera.ttf"))
    pdfmetrics.registerFont(TTFont("VeraBd", "VeraBd.ttf"))

    destination.mkdir(parents=True, exist_ok=True)
    picture = Image.new("RGB", (640, 320), "white")
    draw = ImageDraw.Draw(picture)
    draw.rectangle((160, 70, 360, 255), outline="black", width=6)
    draw.rectangle((168, 150, 352, 248), fill="#b6dbef")
    draw.ellipse((440, 140, 500, 200), fill="#2860c0")
    draw.text((180, 270), "REFERENCE VESSEL", fill="black", font_size=20)
    draw.text((400, 220), "BLUE MARKER", fill="black", font_size=20)
    picture.save(destination / "rich-figure.png")
    pdf = canvas.Canvas(
        str(destination / "rich.pdf"),
        pagesize=(612, 792),
        invariant=1,
        initialFontName="Vera",
    )
    pdf.setTitle("Synthetic field notebook")
    pdf.setAuthor("Tychora synthetic fixtures")

    def text(value, x, y, size=12):
        pdf.setFont("VeraBd" if size > 12 else "Vera", size)
        pdf.drawString(x, 792 - y, value)

    def table(top, rows):
        x, widths, height = 54, (180, 180, 144), 32
        for r, row in enumerate(rows):
            offset = x
            for width, value in zip(widths, row):
                pdf.rect(offset, 792 - top - (r + 1) * height, width, height)
                text(value, offset + 8, top + r * height + 21)
                offset += width

    text(PASSAGES[0], 54, 60, 22)
    text("1. Preparation", 54, 110, 16)
    text(PASSAGES[1], 54, 142)
    text(PASSAGES[2], 54, 169)
    text("Table 1. Observation register", 54, 575, 14)
    table(
        600,
        [
            ("Sample", "Indicator", "Result"),
            ("Aster", "Blue", "Stable"),
            ("Birch", "Green", "Pending"),
        ],
    )
    pdf.showPage()
    text("Table 1. Observation register (continued)", 54, 55, 14)
    table(
        80,
        [
            ("Sample", "Indicator", "Result"),
            ("Cedar", "Amber", "Stable"),
            ("Dahlia", "Violet", "Repeat"),
        ],
    )
    text("2. Observation", 54, 230, 16)
    text(PASSAGES[3], 54, 265)
    pdf.drawImage(ImageReader(picture), 96, 792 - 500, width=420, height=210)
    text(PASSAGES[4], 54, 528)
    pdf.save()

    doc = Document()
    doc.add_heading(PASSAGES[0], 0)
    doc.add_heading("1. Preparation", 1)
    doc.add_paragraph(PASSAGES[1])
    doc.add_paragraph(PASSAGES[2])
    doc.add_paragraph("Table 1. Observation register")
    grid = doc.add_table(rows=1, cols=3)
    grid.style = "Table Grid"
    for cell, value in zip(grid.rows[0].cells, ("Sample", "Indicator", "Result")):
        cell.text = value
    for row in (
        ("Aster", "Blue", "Stable"),
        ("Birch", "Green", "Pending"),
        ("Cedar", "Amber", "Stable"),
        ("Dahlia", "Violet", "Repeat"),
    ):
        for cell, value in zip(grid.add_row().cells, row):
            cell.text = value
    doc.add_page_break()
    doc.add_heading("2. Observation", 1)
    doc.add_paragraph(PASSAGES[3])
    doc.add_picture(str(destination / "rich-figure.png"), width=Inches(5.5))
    doc.add_paragraph(PASSAGES[4], style="Caption")
    doc.save(str(destination / "rich.docx"))


if __name__ == "__main__":
    build(Path(__file__).parent)
